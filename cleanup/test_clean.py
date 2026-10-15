"""Offline tests for the spreadsheet cleaner."""
import io
import json
import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from clean_sheet import (
    clean_date,
    clean_email,
    clean_integer,
    clean_money,
    clean_name,
    clean_phone,
    clean_sku,
    clean_state,
    clean_table,
    main,
    read_table,
    safe_part,
    write_csv,
    write_workbook,
)

HERE = Path(__file__).parent
CONFIG = json.loads((HERE / "customers.config.json").read_text(encoding="utf-8"))
PRODUCTS = json.loads((HERE / "products.config.json").read_text(encoding="utf-8"))


class CleanerTests(unittest.TestCase):
    def test_name(self):
        self.assertEqual(clean_name("  DAVID   O'BRIEN ", {})[0], "David O'Brien")
        self.assertEqual(clean_name("Chen, Wei", {})[0], "Wei Chen")
        self.assertEqual(clean_name("Smith, Jr.", {})[0], "Smith, Jr.")
        self.assertEqual(clean_name("Leslie McDonald", {})[0], "Leslie McDonald")

    def test_email(self):
        self.assertEqual(clean_email(" mailto:Sarah.Lee@Example.com; ", {}), ("sarah.lee@example.com", None))
        self.assertIsNotNone(clean_email("priya.patel@example", {})[1])

    def test_phone(self):
        for raw in ("555.010.1234", "+1 (555) 010-1234", "5550101234", "1-555-010-1234"):
            self.assertEqual(clean_phone(raw, {}), ("(555) 010-1234", None))
        self.assertEqual(clean_phone("555-010-2222 ext. 14", {})[0], "(555) 010-2222 x14")
        self.assertEqual(clean_phone("555-0106", {}), ("555-0106", "not a 10-digit phone number"))

    def test_date(self):
        for raw in ("03/14/2024", "3/14/24", "2024-03-14", "Mar 14, 2024", "March 14th 2024",
                    "14 Mar 2024", "03-14-2024", "2024-03-14 00:00:00"):
            self.assertEqual(clean_date(raw, {}), ("2024-03-14", None), raw)
        self.assertEqual(clean_date("14/03/2024", {"day_first": True})[0], "2024-03-14")
        self.assertIsNotNone(clean_date("2025-02-30", {})[1])

    def test_money(self):
        cases = {"$1,250.00": "1250.00", "1250": "1250.00", "2.4k": "2400.00",
                 "(45.00)": "-45.00", "USD 560": "560.00", "-$12.5": "-12.50"}
        for raw, want in cases.items():
            self.assertEqual(clean_money(raw, {}), (want, None), raw)
        self.assertIsNotNone(clean_money("twelve", {})[1])

    def test_money_accepts_us_thousands_groups(self):
        cases = {"1,234": "1234.00", "1,234,567.89": "1234567.89", "$12,000": "12000.00",
                 "1.5": "1.50", "€45": "45.00", "45€": "45.00", "EUR 560": "560.00",
                 "560 eur": "560.00", "(€1,200.50)": "-1200.50", "-€3k": "-3000.00"}
        for raw, want in cases.items():
            self.assertEqual(clean_money(raw, {}), (want, None), raw)

    def test_money_comma_decimals_go_to_review_by_default(self):
        # Before the fix these were silently read as 1.2005 and 125.00.
        for raw in ("1.200,50", "1,25", "12,3456", "€1,25", "1,2k", "1,23,456"):
            value, problem = clean_money(raw, {})
            self.assertEqual(value, raw, raw)
            self.assertIn("comma decimal?", problem or "", raw)
        self.assertEqual(clean_money("1.200.000", {}),
                         ("1.200.000", "number format unclear (dots as thousands separators?)"))

    def test_money_with_decimal_comma(self):
        opts = {"decimal_comma": True}
        cases = {"1.234,56": "1234.56", "1.200": "1200.00", "45,5": "45.50",
                 "€ 1.234.567,89": "1234567.89", "1.200,50 EUR": "1200.50",
                 "(45,00)": "-45.00", "-12,5": "-12.50", "2,4k": "2400.00", "1250": "1250.00"}
        for raw, want in cases.items():
            self.assertEqual(clean_money(raw, opts), (want, None), raw)

    def test_money_with_decimal_comma_flags_values_that_do_not_fit(self):
        opts = {"decimal_comma": True}
        for raw in ("1,234.56", "1.25", "12.3456", "1,2,3", "$1,250.00"):
            value, problem = clean_money(raw, opts)
            self.assertEqual(value, raw, raw)
            self.assertIn("decimal_comma", problem or "", raw)
        self.assertEqual(clean_money("zwölf", opts)[1], "not a number")

    def test_state(self):
        self.assertEqual(clean_state("california", {})[0], "CA")
        self.assertEqual(clean_state("New  York", {})[0], "NY")
        self.assertEqual(clean_state("tx", {})[0], "TX")
        self.assertIsNotNone(clean_state("Calif.", {})[1])

    def test_sku(self):
        cases = {"  BRK-1042 ": "BRK-1042", '"BRK-1042"': "BRK-1042", "'00412": "00412",
                 "abc-12": "abc-12", "TEE  RED/M": "TEE RED/M", "007.5_A": "007.5_A",
                 "012345678905": "012345678905"}
        for raw, want in cases.items():
            self.assertEqual(clean_sku(raw, {}), (want, None), raw)
        self.assertEqual(clean_sku(" abc-12 ", {"uppercase": True}), ("ABC-12", None))

    def test_sku_needs_review(self):
        for raw in ("BRK#1042", "TEE (RED)", "ab,12", "café-1"):
            value, problem = clean_sku(raw, {"uppercase": True})
            self.assertEqual(value, raw, raw)  # left as is, not uppercased
            self.assertIn("character", problem)
        for raw in ("7.12E+11", "1.23e+11", "4E7"):
            value, problem = clean_sku(raw, {})
            self.assertEqual(value, raw)
            self.assertIn("scientific notation", problem)
            self.assertIn("may have been lost", problem)

    def test_integer(self):
        cases = {"12": "12", " 12 ": "12", "1,200": "1200", "12.0": "12", "0": "0",
                 "1,234,567": "1234567", "1,200.00": "1200"}
        for raw, want in cases.items():
            self.assertEqual(clean_integer(raw, {}), (want, None), raw)

    def test_integer_needs_review(self):
        for raw in ("12 pcs", "about 5", "2.5", "1,20", "12,00", ",120", "+4"):
            value, problem = clean_integer(raw, {})
            self.assertEqual((value, problem), (raw, "not a whole number"), raw)
        value, problem = clean_integer("-3", {})
        self.assertEqual(value, "-3")
        self.assertIn("negative", problem)


class ProductTableTests(unittest.TestCase):
    def test_cleans_a_small_product_table(self):
        table = [
            ["Spring Stock List", "", "", "", ""],
            ["Item #", "Product Name", "Retail", "Qty", "UPC"],
            [" fern-01 ", "USB-C 2PK  cable", "$12.50", "1,200", "'012345678905"],
            ["FERN-02", "linen tote", "USD 30", "12 pcs", "7.12E+11"],
            ["FERN-03", "", "N/A", "-3", ""],
            ["FERN-02", "linen tote", "30", "", ""],
            ["", "", "", "", ""],
            ["Total", "", "$72.50", "", ""],
        ]
        result = clean_table(table, PRODUCTS)
        self.assertEqual(result.header, ["SKU", "Title", "Price", "Quantity", "Barcode"])
        self.assertEqual(result.rows, [
            ["fern-01", "USB-C 2PK cable", "12.50", "1200", "012345678905"],
            ["FERN-02", "linen tote", "30.00", "12 pcs", "7.12E+11"],
            ["FERN-03", "", "", "-3", ""],
        ])
        problems = {(r["source_row"], r["column"], r["problem"].split(";")[0])
                    for r in result.review}
        self.assertEqual(problems, {
            (4, "Quantity", "not a whole number"),
            (4, "Barcode", "looks like a number Excel converted to scientific notation"),
            (5, "Title", "required value is missing"),
            (5, "Price", "required value is missing"),
            (5, "Quantity", "negative quantity"),
        })
        reasons = [r["reason"] for r in result.removed]
        self.assertIn("same SKU as row 4, which has more filled in", reasons)
        self.assertEqual(len(reasons), 4)  # title, duplicate, blank, total

    def test_total_row_needs_the_word_total_not_just_its_letters(self):
        table = [["SKU", "Title", "Price"],
                 ["TT-1", "Totally Tubular Tumbler", "$18.00"],
                 ["TT-2", "Totalizer Pen", "4"],
                 ["Grand Total", "", "$22.00"],
                 ["Sub-total:", "", "$22.00"],
                 ["TOTALS", "", "$22.00"]]
        result = clean_table(table, PRODUCTS)
        self.assertEqual([row[1] for row in result.rows], ["Totally Tubular Tumbler", "Totalizer Pen"])
        self.assertEqual([r["reason"] for r in result.removed], ["total row"] * 3)

    def test_workbook_keeps_flagged_values_as_text(self):
        table = [["SKU", "Title", "Price", "Qty"],
                 ["A-1", "Mug", "$9", "1,200"],
                 ["A-2", "Cup", "twelve", "12 pcs"]]
        result = clean_table(table, PRODUCTS)
        with tempfile.TemporaryDirectory() as tmp:
            write_workbook(result, Path(tmp) / "out.xlsx")
            from openpyxl import load_workbook
            book = load_workbook(Path(tmp) / "out.xlsx")
            rows = list(book["Clean"].iter_rows(min_row=2, values_only=True))
            book.close()
        self.assertEqual(rows, [("A-1", "Mug", 9.0, 1200), ("A-2", "Cup", "twelve", "12 pcs")])


class SampleFileTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.table = read_table(HERE / "sample" / "messy_customers.csv")
        cls.result = clean_table(cls.table, CONFIG)
        cls.by_email = {row[1]: row for row in cls.result.rows}

    def test_header_is_found_below_title_rows_and_renamed(self):
        self.assertEqual(self.result.header, list(CONFIG["columns"]))
        reasons = [r["reason"] for r in self.result.removed]
        self.assertEqual(reasons.count("title row above the header"), 2)

    def test_row_counts_add_up(self):
        r = self.result
        self.assertEqual(len(r.rows), 12)
        self.assertEqual(r.rows_in, len(r.rows) + len(r.removed) - 2)  # 2 title rows sit above the header

    def test_cleans_a_typical_row(self):
        self.assertEqual(self.by_email["wei.chen@example.com"], [
            "Wei Chen", "wei.chen@example.com", "(555) 010-2222 x14", "NY",
            "2025-01-05", "2026-09-12", "89.50", ""])

    def test_duplicates_keep_the_most_complete_row(self):
        self.assertEqual(self.by_email["maria.gonzalez@example.com"][7], "VIP")
        reasons = " / ".join(r["reason"] for r in self.result.removed)
        self.assertIn("same Email as row 5", reasons)
        self.assertIn("exact duplicate of row 12", reasons)

    def test_total_and_blank_rows_are_removed_not_cleaned(self):
        reasons = [r["reason"] for r in self.result.removed]
        self.assertEqual(reasons.count("total row"), 1)
        self.assertEqual(reasons.count("blank row"), 2)
        self.assertFalse(any("TOTAL" in row for row in self.result.rows))

    def test_unfixable_values_are_flagged_and_left_alone(self):
        problems = {(r["column"], r["problem"]) for r in self.result.review}
        self.assertEqual(problems, {
            ("Signup Date", "not a recognizable date"),
            ("Email", "not a valid email address"),
            ("Email", "required value is missing"),
            ("Phone", "not a 10-digit phone number"),
            ("State", "not a US state"),
        })
        self.assertEqual(self.by_email["david.obrien@example.com"][4], "2025-02-30")

    def test_placeholders_become_blank_and_multiline_notes_are_joined(self):
        self.assertEqual(self.by_email["linh.nguyen@example.com"][5:7], ["", ""])
        self.assertEqual(self.by_email["omar.haddad@example.com"][7], "prefers phone calls")

    def test_every_change_is_logged_with_its_original_value(self):
        change = next(c for c in self.result.changes
                      if c["source_row"] == 5 and c["column"] == "Total Spent")
        self.assertEqual((change["before"], change["after"]), ("$1,250.00", "1250.00"))

    def test_unclear_amounts_are_flagged_in_a_table(self):
        table = [["Name", "Total Spent"], ["Ana", "1.200,50"], ["Ben", "1,25"], ["Cy", "$1,250.00"]]
        result = clean_table(table, CONFIG)
        self.assertEqual([row[1] for row in result.rows], ["1.200,50", "1,25", "1250.00"])
        self.assertEqual([(r["source_row"], r["column"]) for r in result.review],
                         [(2, "Total Spent"), (3, "Total Spent")])
        with tempfile.TemporaryDirectory() as tmp:
            write_workbook(result, Path(tmp) / "out.xlsx")  # flagged text stays text
            self.assertEqual([row[1] for row in read_table(Path(tmp) / "out.xlsx")[1:]],
                             ["1.200,50", "1,25", "1250"])

    def test_decimal_comma_is_read_from_the_column_config(self):
        config = {"columns": {"Total": {"type": "money", "decimal_comma": True}}}
        result = clean_table([["Total"], ["1.234,56"], ["45,5"]], config)
        self.assertEqual(result.rows, [["1234.56"], ["45.50"]])
        self.assertEqual(result.review, [])

    def test_outputs_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            write_csv(self.result, Path(tmp) / "out.csv")
            write_workbook(self.result, Path(tmp) / "out.xlsx")
            self.assertEqual(read_table(Path(tmp) / "out.csv")[1:], self.result.rows)
            # An .xlsx input goes through the same pipeline and is already clean.
            again = clean_table(read_table(Path(tmp) / "out.xlsx"), CONFIG)
            self.assertEqual(len(again.rows), len(self.result.rows))
            self.assertEqual(again.removed, [])


HEADER = ["Name", "Email", "Total Spent"]


def make_workbook(path: Path, sheets: dict, hidden: tuple = ()) -> Path:
    """Write a made-up workbook: {sheet name: rows}, in tab order."""
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
        if name in hidden:
            sheet.sheet_state = "hidden"
    book.save(path)
    return path


def run_cli(*args: str) -> tuple[int, str, str]:
    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        code = main(["--config", str(HERE / "customers.config.json"), *args])
    return code, out.getvalue(), err.getvalue()


class SheetTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(self.enterContext(tempfile.TemporaryDirectory()))
        self.book = make_workbook(self.tmp / "export.xlsx", {
            "Summary": [["Report"], ["made up figures"]],
            "Customers": [HEADER, ["ANA LIMA", "ana@example.com", "$10"],
                          ["Ben Ode", "ben@example.com", "20"]],
            "Old List": [HEADER, ["Cy Park", "cy@example.com", "5"]],
            "Blank": [],
        }, hidden=("Old List",))

    def test_read_table_without_sheet_still_reads_the_first(self):
        self.assertEqual(read_table(self.book), [["Report"], ["made up figures"]])

    def test_pick_sheet_by_name(self):
        table = read_table(self.book, sheet="Customers")
        self.assertEqual(table[1], ["ANA LIMA", "ana@example.com", "$10"])

    def test_pick_sheet_ignoring_case_and_spaces(self):
        self.assertEqual(read_table(self.book, sheet="  old list ")[1][0], "Cy Park")

    def test_exact_name_wins_over_a_case_insensitive_match(self):
        path = make_workbook(self.tmp / "two.xlsx", {"data": [["a"]], "Data ": [["b"]]})
        self.assertEqual(read_table(path, sheet="Data "), [["b"]])

    def test_ambiguous_name_is_refused_not_guessed(self):
        path = make_workbook(self.tmp / "two.xlsx", {"data": [["a"]], "Data ": [["b"]]})
        with self.assertRaises(RuntimeError) as caught:
            read_table(path, sheet="DATA")
        self.assertIn("more than one sheet", str(caught.exception))

    def test_unknown_sheet_lists_the_sheets(self):
        with self.assertRaises(RuntimeError) as caught:
            read_table(self.book, sheet="Orders")
        for name in ("Summary", "Customers", "Old List", "Blank"):
            self.assertIn(name, str(caught.exception))

    def test_cli_unknown_sheet_fails_with_the_sheet_list(self):
        code, _, err = run_cli(str(self.book), "--sheet", "Orders", "--out-dir", str(self.tmp))
        self.assertEqual(code, 1)
        self.assertIn('"Customers"', err)

    def test_cli_sheet_with_csv_is_an_error(self):
        code, _, err = run_cli(str(HERE / "sample" / "messy_customers.csv"), "--sheet", "x",
                               "--out-dir", str(self.tmp))
        self.assertEqual(code, 1)
        self.assertIn("Excel", err)
        self.assertEqual(list(self.tmp.glob("*.csv")), [])

    def test_cli_picks_a_sheet_and_names_outputs_after_it(self):
        code, out, _ = run_cli(str(self.book), "--sheet", "customers", "--out-dir", str(self.tmp))
        self.assertEqual(code, 0)
        self.assertIn("2 rows in, 2 rows out", out)
        rows = read_table(self.tmp / "export_Customers_clean.csv")
        self.assertEqual(rows[1], ["Ana Lima", "ana@example.com", "10.00"])

    def test_cli_default_notes_the_sheets_it_did_not_clean(self):
        path = make_workbook(self.tmp / "first.xlsx", {
            "Customers": [HEADER, ["Ana Lima", "ana@example.com", "10"]],
            "Old List": [HEADER, ["Cy Park", "cy@example.com", "5"]],
            "Notes": [["x"]],
        }, hidden=("Old List",))
        code, out, _ = run_cli(str(path), "--out-dir", str(self.tmp))
        self.assertEqual(code, 0)
        self.assertIn('Cleaned sheet "Customers"', out)
        self.assertIn('"Old List" (hidden)', out)
        self.assertIn('"Notes"', out)
        self.assertIn("--sheet", out)
        self.assertTrue((self.tmp / "first_clean.csv").exists())

    def test_cli_single_sheet_workbook_has_no_notice(self):
        path = make_workbook(self.tmp / "one.xlsx", {
            "Customers": [HEADER, ["Ana Lima", "ana@example.com", "10"]]})
        code, out, _ = run_cli(str(path), "--out-dir", str(self.tmp))
        self.assertEqual(code, 0)
        self.assertNotIn("--sheet", out)

    def test_cli_all_cleans_each_sheet_and_skips_empty_ones(self):
        code, out, _ = run_cli(str(self.book), "--sheet", "all", "--out-dir", str(self.tmp))
        self.assertEqual(code, 0)
        written = sorted(p.name for p in self.tmp.glob("export_*_clean.*"))
        self.assertEqual(written, [
            "export_Customers_clean.csv", "export_Customers_clean.xlsx",
            "export_Old_List_clean.csv", "export_Old_List_clean.xlsx",
            "export_Summary_clean.csv", "export_Summary_clean.xlsx",
        ])
        self.assertIn('Sheet "Blank": no data rows, skipped', out)
        self.assertIn('Sheet "Old List" (hidden)', out)
        # Sheets are cleaned separately: Old List's row is not merged into Customers.
        self.assertEqual(len(read_table(self.tmp / "export_Customers_clean.csv")), 3)
        self.assertEqual(len(read_table(self.tmp / "export_Old_List_clean.csv")), 2)

    def test_cli_all_fails_when_every_sheet_is_empty(self):
        path = make_workbook(self.tmp / "empty.xlsx", {"A": [], "B": [["Name", "Email"]]})
        code, _, err = run_cli(str(path), "--sheet", "all", "--out-dir", str(self.tmp))
        self.assertEqual(code, 1)
        self.assertIn("no data rows", err)

    def test_cli_sheet_literally_named_all(self):
        path = make_workbook(self.tmp / "named.xlsx", {
            "Other": [HEADER, ["Cy Park", "cy@example.com", "5"]],
            "All": [HEADER, ["Ana Lima", "ana@example.com", "10"]]})
        code, _, _ = run_cli(str(path), "--sheet", "=all", "--out-dir", str(self.tmp))
        self.assertEqual(code, 0)
        self.assertEqual(sorted(p.name for p in self.tmp.glob("*.csv")), ["named_All_clean.csv"])

    def test_sheet_names_are_made_safe_and_kept_distinct(self):
        self.assertEqual(safe_part('Q1 "Sales" <draft>'), "Q1_Sales_draft")
        self.assertEqual(safe_part("..."), "sheet")
        path = make_workbook(self.tmp / "clash.xlsx", {
            "Q1 Sales": [HEADER, ["Ana Lima", "ana@example.com", "10"]],
            "Q1_Sales": [HEADER, ["Cy Park", "cy@example.com", "5"]]})
        code, _, _ = run_cli(str(path), "--sheet", "all", "--out-dir", str(self.tmp))
        self.assertEqual(code, 0)
        self.assertEqual(sorted(p.name for p in self.tmp.glob("*.csv")),
                         ["clash_Q1_Sales_2_clean.csv", "clash_Q1_Sales_clean.csv"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
